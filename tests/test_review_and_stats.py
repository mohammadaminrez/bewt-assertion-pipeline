from __future__ import annotations

import pytest

from src.data.store import ResultStore
from src.evaluation.comparator import answer_key, classify_outcome
from src.evaluation.excel_export import export_review_to_excel, import_review_labels, review_groups
from src.evaluation.metrics import _cochrans_q, _mcnemar_exact, _wilcoxon, holm, statistical_tests
from src.evaluation.reporter import generate_full_report
from src.models import ErrorCategory, ExperimentResult, TestRecord
from src.parsing.java_parser import extract_assertions

GOLD = 'assertEquals("Create your own content", page.getFirstArticle());'


def _record(name="SeeCategoryArticlesTest", gold=GOLD):
    return TestRecord(app="bludit", variant="baseline", version="3.13.1", file_path="", class_name=name,
                      method_name="m", assertions=extract_assertions(gold))


def _result(treatment, run, answer, passes=True, compiles=True, model="gpt-4o-mini", name="SeeCategoryArticlesTest"):
    rec = _record(name)
    category = classify_outcome(answer, rec.assertions, compiles, passes, "" if passes else "java.lang.AssertionError")
    return ExperimentResult(test_record=rec, treatment=treatment, model=model, run=run, prompt="p",
                            raw_response=answer, generated_assertion=answer, compiles=compiles,
                            passes=passes, error_category=category)


def test_outcome_categories_are_decided_only_when_certain():
    gold = _record().assertions
    assert classify_outcome(GOLD, gold, True, True) is ErrorCategory.CORRECT
    assert classify_outcome("  " + GOLD.replace(", ", ",\n   "), gold, True, True) is ErrorCategory.CORRECT
    assert classify_outcome('assertTrue(page.x().contains("Create"));', gold, True, True) is ErrorCategory.NEEDS_REVIEW
    assert classify_outcome(GOLD, gold, False, False) is ErrorCategory.NOT_EXECUTABLE
    assert classify_outcome(GOLD, gold, True, False, "org.junit.ComparisonFailure") is ErrorCategory.WRONG_ASSERTION
    assert classify_outcome(GOLD, gold, True, False, "org.openqa.selenium.NoSuchElementException") \
        is ErrorCategory.NOT_EXECUTABLE


def test_same_answer_in_several_runs_is_reviewed_once_and_label_applies_to_all(tmp_path):
    weak = 'assertTrue(page.firstArticle.getText().contains("Create your own content"));'
    results = [_result("D", 1, weak), _result("D", 2, "  " + weak), _result("C", 1, weak),
               _result("A", 1, GOLD), _result("B", 1, 'assertTrue(page.isShown());')]
    groups = review_groups(results)
    assert [len(g["occurrences"]) for g in groups] == [3, 1]

    store = ResultStore(tmp_path / "results.db")
    for r in results:
        store.save_result(r)
    path = tmp_path / "review.xlsx"
    export_review_to_excel(review_groups(store.load_experiment_results()), path)

    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["Review"]
    header = [c.value for c in ws[1]]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, header.index("Generated Assertion (passed)") + 1).value.startswith("assertTrue(page.first"):
            ws.cell(row, header.index("Manual Classification") + 1).value = "under_assertive"
    wb.save(path)

    labels = import_review_labels(path)
    assert len(labels) == 1
    updated = sum(store.apply_review_label(l["app"], l["class_name"], l["answer_key"], l["label"], l["notes"])
                  for l in labels)
    assert updated == 3
    final = sorted((r.treatment, r.run, r.final_category) for r in store.load_experiment_results())
    assert final == [("A", 1, "correct"), ("B", 1, "needs_review"), ("C", 1, "under_assertive"),
                     ("D", 1, "under_assertive"), ("D", 2, "under_assertive")]
    assert len(review_groups(store.load_experiment_results())) == 1
    store.close()


def test_invalid_review_label_is_rejected(tmp_path):
    path = tmp_path / "review.xlsx"
    export_review_to_excel(review_groups([_result("D", 1, "assertTrue(page.ok());")]), path)
    import openpyxl
    wb = openpyxl.load_workbook(path)
    wb["Review"].cell(2, 9).value = "not_executable"
    wb.save(path)
    with pytest.raises(ValueError, match="Invalid labels"):
        import_review_labels(path)


def test_answer_key_ignores_whitespace_and_markers():
    assert answer_key("// (1)\nassertTrue(a.b());") == answer_key("assertTrue( a.b() );".replace("( ", "(").replace(" )", ")"))
    assert answer_key("assertTrue(a.b());") != answer_key("assertFalse(a.b());")


def test_holm_adjustment():
    assert holm({"a": 0.01, "b": 0.04, "c": 0.03}) == pytest.approx({"a": 0.03, "b": 0.06, "c": 0.06})


def test_cochrans_q_and_mcnemar_known_values():
    table = [[1, 1, 1], [1, 1, 0], [1, 0, 0], [0, 0, 0], [1, 1, 0], [1, 0, 0],
             [1, 1, 0], [1, 0, 0], [0, 0, 0], [1, 1, 1], [1, 0, 0], [1, 1, 0]]
    # independent form: Q = k(k-1) * sum((C_j - mean C)^2) / (k*N - sum(R_i^2))
    k, cols = 3, [sum(r[j] for r in table) for j in range(3)]
    rows, n = [sum(r) for r in table], sum(sum(r) for r in table)
    expected = k * (k - 1) * sum((c - n / k) ** 2 for c in cols) / (k * n - sum(r * r for r in rows))
    q = _cochrans_q(table)
    assert q["Q"] == pytest.approx(expected) == pytest.approx(12.0) and q["df"] == 2
    assert _mcnemar_exact([1, 1, 1, 1, 0], [0, 0, 0, 0, 0]) == {"x_only": 4, "y_only": 0, "p": pytest.approx(0.125)}
    assert _mcnemar_exact([1, 0], [1, 0])["p"] == 1.0


def test_wilcoxon_handles_identical_samples_and_direction():
    assert _wilcoxon([0.2, 0.4], [0.2, 0.4]) == {"statistic": None, "p": 1.0, "n_nonzero": 0, "rank_biserial": 0.0}
    res = _wilcoxon([1, 1, 1, 1, 1, 0.8], [0, 0.2, 0, 0, 0.4, 0.8])
    assert res["n_nonzero"] == 5 and res["rank_biserial"] == 1.0


def test_report_uses_runs_and_final_categories(tmp_path):
    results = []
    for name, d_pass in (("T1", True), ("T2", True), ("T3", False)):
        for run in (1, 2, 3):
            results.append(_result("A", run, "assertTrue(x.y());", passes=False, name=name))
            results.append(_result("D", run, GOLD, passes=d_pass or run == 1, name=name))
    stats = statistical_tests(results)
    model = stats["models"]["gpt-4o-mini"]
    assert model["tests"] == 3 and model["treatments"] == ["A", "D"]
    assert model["primary_pass_rate"]["pairwise_wilcoxon"]["A_vs_D"]["rank_biserial"] == -1.0
    generate_full_report(results, tmp_path)
    rows = (tmp_path / "by_treatment.csv").read_text().splitlines()
    assert rows[0].startswith("model,treatment,tests,runs,answers,pass_rate,pass_rate_sd_runs")
    d_row = next(r for r in rows if r.startswith("gpt-4o-mini,D,"))
    assert d_row.split(",")[2:6] == ["3", "3", "9", "0.7778"]
    assert "pass rate on the live application" in (tmp_path / "latex_tables.tex").read_text()
